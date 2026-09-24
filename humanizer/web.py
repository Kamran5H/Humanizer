"""
Humanizer Pro Web Studio Server.
High-performance FastAPI web application providing real-time AI detection analysis,
stealth document humanization, visual diff generation, and run-preserved DOCX export.
"""

from __future__ import annotations

import logging
import os
import re
import shutil
import tempfile
import time
import uuid
import webbrowser
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse
from pydantic import BaseModel

from humanizer.config import HumanizerConfig
from humanizer.detector import score_text
from humanizer.document import extract_text_from_file, extract_text_from_pdf, iter_document_paragraphs
from humanizer.engine import humanize_docx, humanize_text
from humanizer.lint import lint_text, readability
from humanizer.providers import get_provider_status, reset_cooldowns

logger = logging.getLogger(__name__)

# Temporary upload session storage
_UPLOAD_DIR = Path(tempfile.gettempdir()) / "humanizer_pro_uploads"
_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
_SESSION_ID_RE = re.compile(r"^[a-zA-Z0-9_-]{1,64}$")

app = FastAPI(
    title="Humanizer Pro Studio",
    description="Academic & Journal-Grade Stealth Humanization Studio",
    version="5.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

_TEMPLATES_DIR = Path(__file__).resolve().parent / "templates"
_INDEX_HTML_PATH = _TEMPLATES_DIR / "index.html"


class ScoreRequest(BaseModel):
    text: str


class HumanizeRequest(BaseModel):
    text: str
    scientific: bool = True
    stealth_level: int = 3
    target_pct: float = 5.0
    use_gemini: bool = True
    session_docx: Optional[str] = None


class DownloadDocxRequest(BaseModel):
    text: str
    session_docx: Optional[str] = None


@app.get("/", response_class=HTMLResponse)
async def serve_index() -> HTMLResponse:
    """Serve the single-page Web Studio interface."""
    if not _INDEX_HTML_PATH.exists():
        raise HTTPException(status_code=404, detail="Studio template not found")
    content = _INDEX_HTML_PATH.read_text(encoding="utf-8")
    return HTMLResponse(content=content)


@app.post("/api/score")
async def api_score(req: ScoreRequest) -> dict:
    """Analyze text with multi-signal AI detector and readability metrics."""
    text = (req.text or "").strip()
    if not text:
        return {
            "ai_score": 0.0,
            "perplexity": 0.0,
            "burstiness": 0.0,
            "smooth_frac": 0.0,
            "ai_tells": 0,
            "sentences": 0,
            "ttr": 1.0,
            "flesch": None,
            "words": 0,
            "chars": 0,
        }

    scores = score_text(text)
    read = readability(text)
    words = len(text.split())

    return {
        **scores,
        "flesch": read.get("flesch"),
        "words": words,
        "chars": len(text),
        "lint_issues": lint_text(text)[:5],
    }


@app.post("/api/humanize")
async def api_humanize(req: HumanizeRequest) -> dict:
    """Execute stealth humanization across the input document."""
    text = (req.text or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="Empty text provided")

    t0 = time.time()
    init_res = score_text(text)

    cfg = HumanizerConfig(
        use_gemini=req.use_gemini,
        stealth_level=req.stealth_level,
        target_pct=req.target_pct,
        scientific=req.scientific,
    )

    try:
        # If an existing uploaded DOCX file exists, process it with run-preservation
        if req.session_docx:
            if not _SESSION_ID_RE.match(req.session_docx):
                raise HTTPException(status_code=400, detail="Invalid session ID format")
            src_doc = _UPLOAD_DIR / f"{req.session_docx}.docx"
            if src_doc.exists():
                out_doc = _UPLOAD_DIR / f"{req.session_docx}_humanized.docx"
                saved = humanize_docx(src_doc, cfg, dst=out_doc)
                import docx
                d = docx.Document(str(saved))
                out_text = "\n\n".join(p.text for p in iter_document_paragraphs(d) if p.text.strip())
            else:
                out_text = humanize_text(text, cfg)
        else:
            out_text = humanize_text(text, cfg)

        final_res = score_text(out_text)
        read = readability(out_text)
        final_res["flesch"] = read.get("flesch")
        dur = time.time() - t0

        return {
            "success": True,
            "output": out_text,
            "initial_score": init_res.get("ai_score", 0.0),
            "final_score": final_res.get("ai_score", 0.0),
            "metrics": final_res,
            "duration": round(dur, 2),
        }
    except Exception as e:
        logger.error(f"Humanization failed: {e}", exc_info=True)
        return {
            "success": False,
            "error": str(e),
        }


@app.post("/api/upload")
async def api_upload(file: UploadFile = File(...)) -> dict:
    """Upload and extract clean text from .docx, .pdf, or .txt documents."""
    filename = file.filename or "document"
    suffix = Path(filename).suffix.lower()

    if suffix not in (".docx", ".pdf", ".txt", ".md"):
        raise HTTPException(status_code=400, detail="Supported file formats: .docx, .pdf, .txt, .md")

    session_id = str(uuid.uuid4())[:8]
    temp_path = _UPLOAD_DIR / f"{session_id}{suffix}"

    with open(temp_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    try:
        extracted_text = extract_text_from_file(temp_path)
        words = len(extracted_text.split())
        paras = len([p for p in extracted_text.split("\n\n") if p.strip()])

        return {
            "success": True,
            "filename": filename,
            "text": extracted_text,
            "words": words,
            "paragraphs": paras,
            "session_docx": session_id if suffix == ".docx" else None,
        }
    except Exception as e:
        logger.error(f"Upload extraction failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Failed to extract text from file: {e}")


@app.post("/api/download-docx")
async def api_download_docx(req: DownloadDocxRequest):
    """Generate and return a cleanly formatted Word (.docx) document."""
    text = (req.text or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="Empty text")

    # 1. If we have a processed humanized DOCX file from session, return it directly
    if req.session_docx:
        if not _SESSION_ID_RE.match(req.session_docx):
            raise HTTPException(status_code=400, detail="Invalid session ID format")
        out_doc = _UPLOAD_DIR / f"{req.session_docx}_humanized.docx"
        if out_doc.exists():
            return FileResponse(
                path=str(out_doc),
                filename="humanized_document.docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    # 2. Otherwise synthesize a clean structured DOCX from the output text
    import docx
    doc = docx.Document()
    for block in text.split("\n\n"):
        if block.strip():
            doc.add_paragraph(block.strip())

    out_temp = _UPLOAD_DIR / f"export_{uuid.uuid4().hex[:8]}.docx"
    doc.save(str(out_temp))

    return FileResponse(
        path=str(out_temp),
        filename="humanized_document.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.get("/api/providers")
async def api_providers() -> list[dict]:
    """Return live status and latency of the 14-endpoint provider pool."""
    return get_provider_status()


@app.post("/api/reset-cooldowns")
async def api_reset_cooldowns() -> dict:
    """Clear all rate-limit cooldowns."""
    reset_cooldowns()
    return {"success": True}


def launch_web(host: str = "127.0.0.1", port: int = 8000, open_browser: bool = True) -> None:
    """Boot the Uvicorn server and automatically open the studio in the default browser."""
    import threading
    import uvicorn

    def _open_url():
        time.sleep(1.2)
        webbrowser.open(f"http://{host}:{port}")

    if open_browser:
        threading.Thread(target=_open_url, daemon=True).start()

    print(f"\n🚀 Humanizer Pro Studio running at: http://{host}:{port}\n")
    uvicorn.run(app, host=host, port=port, log_level="info")


if __name__ == "__main__":
    launch_web()
