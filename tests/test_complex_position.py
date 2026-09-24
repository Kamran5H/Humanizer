import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from humanizer.document import _tag_runs, _reconstruct_runs, _extract_run_style

doc = docx.Document()
p = doc.add_paragraph()
r1 = p.add_run("The relationship ")
# Add a run with an inline drawing/math
r_math = p.add_run()
drawing = OxmlElement('w:drawing')
r_math._element.append(drawing)
r2 = p.add_run(" explains the energy output.")

print("Original runs count:", len(p.runs))
tagged_text, vault = _tag_runs(p)
print("Tagged text:", repr(tagged_text))
print("Vault keys:", list(vault.keys()))

base_style = _extract_run_style(p.runs[0])
rewritten = tagged_text.replace("The relationship", "This known relationship")

_reconstruct_runs(p, rewritten, vault, base_style)

print("Reconstructed runs count:", len(p.runs))
for i, r in enumerate(p.runs):
    has_drawing = "<w:drawing" in r._element.xml
    print(f"  Run {i}: text={r.text!r}, has_drawing={has_drawing}")

# The drawing run should be between "This known relationship " and " explains the energy output."!
# NOT at the end!
runs_with_text = [r for r in p.runs if r.text]
assert "<w:drawing" in p.runs[-1]._element.xml, "Drawing is at end or where is it?"
print("Drawing run position in p.runs: last run index =", len(p.runs)-1)
# Check if it was placed after the trailing text:
last_run_text = p.runs[-2].text if len(p.runs) >= 2 else ""
print(f"Run before drawing: text={last_run_text!r}")
if "explains the energy output" in last_run_text:
    print("CRITICAL DEFECT CONFIRMED: Drawing was moved to the very END of the paragraph after trailing text!")
