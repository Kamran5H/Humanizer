import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
from humanizer.document import _tag_runs, _reconstruct_runs, _extract_run_style, RunStyle

doc = docx.Document()
p = doc.add_paragraph()
r1 = p.add_run("Zinc-air batteries are promising")
r2 = p.add_run("1,2")
r2.font.superscript = True
r3 = p.add_run(". However, degradation occurs.")

print("Original paragraph text:", repr(p.text))
print("Original runs:")
for i, r in enumerate(p.runs):
    print(f"  Run {i}: text={r.text!r}, superscript={r.font.superscript if r.font else None}")

tagged_text, vault = _tag_runs(p)
print("\nTagged text:", repr(tagged_text))
print("Vault:", vault)

base_style = _extract_run_style(p.runs[0])
# Simulate rewrite that kept the citation placeholder or text
rewritten = tagged_text.replace("Zinc-air batteries are promising", "Aqueous zinc-air systems show promise")
print("\nRewritten text to reconstruct:", repr(rewritten))

_reconstruct_runs(p, rewritten, vault, base_style)

print("\nReconstructed paragraph text:", repr(p.text))
print("Reconstructed runs:")
for i, r in enumerate(p.runs):
    print(f"  Run {i}: text={r.text!r}, superscript={r.font.superscript if r.font else None}")

# Check if r2 is still superscript!
cite_run = [r for r in p.runs if "1,2" in r.text]
assert len(cite_run) > 0, "Citation run missing!"
assert cite_run[0].font and cite_run[0].font.superscript is True, f"SUPERSCRIPT WAS LOST! Got {cite_run[0].font.superscript if cite_run[0].font else None}"
print("\nSUCCESS: Superscript preserved!")
